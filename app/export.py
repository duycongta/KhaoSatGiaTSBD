import io
from openpyxl import Workbook
from openpyxl.styles import Font

def to_xlsx(query: dict, listings, res) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Tổng hợp"
    ws.append(["KHẢO GIÁ THAM KHẢO - không phải kết quả thẩm định giá"]); ws["A1"].font = Font(bold=True)
    for k, v in query.items(): ws.append([k, v])
    ws.append([])
    ws.append(["Số tin dùng / tổng", f"{res.n_used}/{res.n_total}"])
    for lbl, v in (("Trung vị đ/m²", res.median), ("P25 đ/m²", res.p25), ("P75 đ/m²", res.p75),
                   ("Khoảng giá sau chiết khấu - thấp", res.low), ("Khoảng giá sau chiết khấu - cao", res.high)):
        ws.append([lbl, round(v) if v else None])
    ws.append(["Cảnh báo", res.warning])
    w2 = wb.create_sheet("Tin so sánh")
    w2.append(["Nguồn", "Link", "Tiêu đề", "Địa chỉ", "Loại đất", "Diện tích m²", "Mặt tiền m", "Giá (đ)", "Giá/m²", "Ngày đăng", "Thu thập lúc", "Loại trừ (ngoại lai)"])
    for c in w2[1]: c.font = Font(bold=True)
    out_ids = {id(x) for x in res.outliers}
    for l in listings:
        w2.append([l.source, l.url, l.title, l.address, l.land_type, l.area_m2, l.frontage_m,
                   l.price_vnd, round(l.price_per_m2) if l.price_per_m2 else None,
                   l.posted_at, l.fetched_at, "x" if id(l) in out_ids else ""])
    b = io.BytesIO(); wb.save(b); return b.getvalue()
